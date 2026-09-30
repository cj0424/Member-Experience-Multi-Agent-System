"""
Execution Kit Agent — takes ONE approved recommendation and reasons
freely about what it needs to become real. Two output types,
matched to real-world use:
- Communications (staff messages, member announcements, internal
  protocols, vendor requests) stay as TEXT — covers the large
  majority of real pieces, and the club can copy/paste, email, or
  print it themselves if needed.
- Tracking/verification logs become a REAL .xlsx file (structured,
  formatted, ready to fill in or upload to Google Sheets for
  real-time shared use — this system builds the file, the club
  handles the optional upload step).

LIMITATIONS (by design):
- Drafts text/files only. Never sends, posts, purchases, uploads,
  or implements anything.
- Cannot know real budget, staffing, or calendar.
- Cannot make the decision — that already happened in Action
  Planning's approval step.
- Cannot verify anything got done — that's Outcome Check's job.

Every prompt receives the shared RAG library (rag.py): the club profile
(staff, shifts, channels, tools, rules), so each piece fits how this club
actually works, plus the technical references for pieces that need them
(e.g. a maintenance protocol).

Phase 3 rules (after the first real kit, where a Google reply said the
fix was already done and promised diffusers the plan only buys if needed):
- public and copy-ready messages never say an action is done before it is;
  they're marked "para publicar/enviar cuando…"
- optional parts of the plan are never presented as certain
- verification: at least two checks on different days, including the
  person who reported the problem if it came from staff
- style: instructions may be impersonal ("se recomienda", "conviene");
  copy-ready texts to a person or group are direct and natural
"""

import json
from rag import load_rag_library, load_club_profile
from llm import generate
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SHARED_QUALITY_RULES = """Antes de incluir CUALQUIER pieza, aplica esta prueba estricta:
¿esto le da al lector algo que genuinamente NO sabría o NO tendría
ya preparado por sí mismo? Si estás a punto de explicar un
procedimiento básico que cualquier persona en ese puesto ya
conocería, DESCÁRTALA — no aporta valor real, sin importar el tema.

No existe un número fijo de piezas — depende genuinamente de lo
que esta recomendación necesite. Cada pieza debe pasar la prueba
de valor real; descarta cualquier pieza de relleno.

Para cualquier texto dirigido directamente a un socio o cliente,
evita un tono presuntuoso que le diga cómo debería sentirse (ej.
"esperamos que notes la diferencia") — informa de forma natural y
directa, sin generar una expectativa que luego tengas que cumplir.

Esto también aplica a instrucciones internas dirigidas al personal
o al propietario: evita un tono de mandato rígido ("debe", "tiene
que", "está obligado a") cuando estés sugiriendo un método de
verificación o una buena práctica de gestión — usa un tono de
recomendación profesional en su lugar ("se recomienda que...",
"lo ideal es que...", "conviene que..."). El contenido puede ser
igual de específico y firme; solo cambia el tono de orden a
sugerencia experta.

IMPORTANTE — no asumas datos ni infraestructura que el sistema no
puede confirmar que existan:

1. Si una pieza asume que el club puede identificar o contactar
   directamente a socios específicos, NO redactes un mensaje
   dirigido a esa persona — en su lugar, redacta un mensaje GENERAL
   a todos los socios, más una nota aparte para el propietario
   sugiriendo contacto directo SI tiene forma de identificarlos.

2. Si una pieza asume que el club ya tiene una herramienta
   específica configurada (WhatsApp Business, un sistema de
   tickets), acláralo explícitamente en vez de asumir que ya existe.

3. Cuando el kit incluya un método de verificación físico o manual,
   sé ESPECÍFICO: indica quién exactamente lo hace (rol concreto,
   no solo "el personal") y en qué momento (apertura, cierre, etc.).

IMPORTANTE — no des nada por hecho antes de que ocurra:

4. Los mensajes públicos o listos para enviar (respuestas a reseñas
   de Google, avisos a socios, mensajes al grupo del personal,
   correos) NUNCA dicen que una acción ya está hecha si el kit se
   prepara antes de hacerla. Escríbelos para enviarlos DESPUÉS, e
   indícalo en el título de la pieza (por ejemplo, "para publicar
   cuando el ajuste esté comprobado"). Si el mensaje se envía antes
   de la acción, redáctalo en futuro ("vamos a revisar…"), sin
   prometer fechas ni resultados que el club no controla.

5. Lo que la recomendación marca como opcional o condicional (por
   ejemplo, "si hace falta, comprar difusores") NUNCA aparece como
   seguro en ningún mensaje. En las instrucciones internas, déjalo
   como condición ("si tras el ajuste…"); en los mensajes públicos,
   menciona solo lo que la recomendación hace con seguridad.

6. En "Cómo verificar", pide al menos DOS comprobaciones en días
   distintos (una sola comprobación puede ser casualidad). Si la
   recomendación indica que el problema lo detectó alguien del
   personal (por ejemplo, un entrenador), que esa persona participe
   en la comprobación. Indica que el resultado de cada comprobación
   se anota en las OBSERVACIONES DEL PERSONAL (no en la nota de
   traspaso ni en el parte de incidencias): es lo que el sistema lee
   para saber si el plan funciona.

FORMATO DE SALIDA — dos tipos, según lo que la pieza realmente sea:
- Comunicaciones (mensajes, correos, avisos, protocolos internos,
  información de referencia) van como TEXTO NORMAL dentro del kit
  — esto cubre la gran mayoría de piezas.
- Si una pieza es un REGISTRO DE SEGUIMIENTO/TRACKER (una tabla que
  alguien rellena repetidamente con fecha y firma), no la escribas
  solo como texto — añade, al final de tu respuesta, un bloque:
  ===TRACKER_SPEC===
  {"titulo": "...", "columnas": ["...", "..."], "fila_ejemplo": ["...", "..."]}
  ===FIN_TRACKER_SPEC===
- Si no hay ningún tracker, omite este bloque por completo. Ante la
  duda entre tracker y texto normal, prefiere texto normal.
- ANTES de generar cualquier pieza nueva, revisa las piezas que ya
  escribiste en este mismo kit — si la nueva pieza repetiría
  información ya cubierta en otra, NO la generes por separado;
  intégrala en la pieza existente o descártala.
- Escribe cada texto listo para enviar en líneas que empiecen con ">".
- Pon en **negrita** solo la acción clave de cada instrucción (3-6
  palabras), y como máximo 3 negritas por pieza. Nunca pongas en
  negrita una frase completa.
- Las líneas "Ojo con" de la recomendación son ADVERTENCIAS, no
  acciones aprobadas: tenlas en cuenta al redactar, pero nunca las
  conviertas en una pieza propia.
- Usa el PERFIL DEL CLUB para ajustar cada pieza a cómo funciona
  realmente este club (personal y turnos, canales, herramientas,
  normas, quién aprueba qué). Lo marcado "por confirmar" no lo des por
  hecho: indícalo como algo a comprobar.
- Usa la BIBLIOTECA DE REFERENCIA solo si una pieza la necesita (por
  ejemplo, un protocolo técnico de mantenimiento).
- Escribe en español de España, tanto en los mensajes a socios como
  en las instrucciones al personal. Usa con el propietario y con los
  socios el tratamiento que indica el PERFIL DEL CLUB (tú o usted).

ESTILO:
- Recomendaciones e instrucciones: pueden ser impersonales ("se
  recomienda", "conviene", "se informará"). Es lo normal.
- Textos listos para enviar a una persona o a un grupo (respuestas
  en Google, correos, WhatsApp, avisos a socios): directos y
  naturales, como los escribiría una persona del club.
- Evita las fórmulas burocráticas en cualquier texto: "proceder a"
  ("hemos procedido a ajustar" → "hemos ajustado"), "llevar a cabo",
  "a la mayor brevedad", "en aras de"."""


def build_prompt(approved_recommendation: str) -> str:
    return f"""Eres un asistente que ayuda al propietario de un club de pádel
a convertir una recomendación YA APROBADA en piezas reales y listas
para usar.

RECOMENDACIÓN APROBADA:
{approved_recommendation}

PERFIL DEL CLUB:
{load_club_profile()}

BIBLIOTECA DE REFERENCIA:
{load_rag_library()}

Antes de redactar nada, analiza la recomendación y pregúntate:
- ¿Hay algo que comunicar a alguien? Redacta el texto exacto.
- ¿Hay algo físico que adquirir o preparar? Indícalo con claridad.
- ¿Hay algo que requiere contactar a personas específicas?
- ¿Hay algún registro de seguimiento que necesite un tracker real?
- ¿Cada mensaje se envía antes o después de la acción? Redáctalo
  según eso (regla 4).

{SHARED_QUALITY_RULES}

Reglas adicionales:
- NO inventes acciones nuevas — trabaja solo con lo ya aprobado.
- Tu única función es REDACTAR. Nunca envíes, publiques, compres
  ni implementes nada.

Formato:

📄 KIT DE EJECUCIÓN

[Piezas que pasen la prueba de calidad, con el bloque TRACKER_SPEC
donde corresponda]

📋 Ejecución:
Responsable: [rol]
Plazo sugerido: [plazo]
Cómo verificar: [específico — quién, cuándo, y al menos dos comprobaciones en días distintos]

"Cómo verificar" va siempre al final, en la MISMA línea que la etiqueta (si hay
varias comprobaciones: "(1) … (2) …"), no como una lista debajo.

Ahora razona desde cero y redacta el kit de ejecución real."""


def build_revision_prompt(approved_recommendation: str, previous_kit: str, owner_feedback: str) -> str:
    return f"""Eres un asistente que ayuda al propietario de un club de pádel.
Ya generaste un kit, pero el propietario ha pedido cambios.

RECOMENDACIÓN APROBADA:
{approved_recommendation}

PERFIL DEL CLUB:
{load_club_profile()}

BIBLIOTECA DE REFERENCIA:
{load_rag_library()}

KIT ANTERIOR:
{previous_kit}

FEEDBACK DEL PROPIETARIO:
{owner_feedback}

Genera un kit NUEVO que responda directamente al feedback.

{SHARED_QUALITY_RULES}

Antes del kit revisado, añade:
"✏️ Cambios respecto a tu feedback: [1 frase]"

Formato:

📄 KIT DE EJECUCIÓN

[Piezas, con bloque TRACKER_SPEC donde corresponda]

📋 Ejecución:
Responsable: [rol]
Plazo sugerido: [plazo]
Cómo verificar: [específico — quién, cuándo, y al menos dos comprobaciones en días distintos]

"Cómo verificar" va siempre al final, en la MISMA línea que la etiqueta (si hay
varias comprobaciones: "(1) … (2) …"), no como una lista debajo.

Ahora redacta el kit revisado."""


def extract_tracker_specs(kit_text: str) -> list[dict]:
    """Pulls out every TRACKER_SPEC block present."""
    specs = []
    parts = kit_text.split("===TRACKER_SPEC===")[1:]
    for part in parts:
        try:
            block = part.split("===FIN_TRACKER_SPEC===")[0].strip()
            specs.append(json.loads(block))
        except (IndexError, json.JSONDecodeError):
            continue
    return specs


def build_tracker_excel(spec: dict, output_path: str):
    """Turns a tracker spec into a real, formatted .xlsx file —
    ready to fill in, print, or upload to Google Sheets for
    real-time shared use."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registro"

    n_cols = len(spec["columnas"])
    ws["A1"] = spec["titulo"]
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin = Border(*[Side(style="thin")] * 4)

    for col, h in enumerate(spec["columnas"], start=1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin

    for col, val in enumerate(spec.get("fila_ejemplo", []), start=1):
        cell = ws.cell(row=4, column=col, value=val)
        cell.font = Font(name="Arial", italic=True, color="888888")
        cell.border = thin

    for row in range(5, 15):
        for col in range(1, n_cols + 1):
            ws.cell(row=row, column=col, value="").border = thin

    for col in range(1, n_cols + 1):
        ws.column_dimensions[chr(64 + col)].width = 16

    wb.save(output_path)


def safe_filename(title: str) -> str:
    """Turns a tracker's real title into a clean, valid filename —
    used instead of a generic counter so files stay identifiable."""
    safe = "".join(c if c.isalnum() or c in " -_" else "" for c in title)
    return safe.strip().replace(" ", "_")


def run_execution_kit_agent(approved_recommendation: str):
    prompt = build_prompt(approved_recommendation)
    return generate(prompt, agent="execution_kit")


def run_execution_kit_revision(approved_recommendation: str, previous_kit: str, owner_feedback: str):
    prompt = build_revision_prompt(approved_recommendation, previous_kit, owner_feedback)
    return generate(prompt, agent="execution_kit")


if __name__ == "__main__":
    test_approved_recommendation = """Protocolo interno de cepillado regular (Coste mínimo): Establecer
una rutina fija para que el personal de pista pase el cepillo ancho
o la rastra 2 o 3 veces por semana, siempre en seco, con un
registro de seguimiento firmado."""

    kit = run_execution_kit_agent(test_approved_recommendation)
    print(kit)

    trackers = extract_tracker_specs(kit)
    for spec in trackers:
        path = f"outputs/{safe_filename(spec['titulo'])}.xlsx"
        build_tracker_excel(spec, path)
        print(f"\n✅ Tracker Excel generado: {path}")
